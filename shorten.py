from colorama import Fore, Style, init
import docx
init()
class line:
    BOLD = Style.BRIGHT
    BLUE = Fore.BLUE
    CYAN = Fore.CYAN
    RED = Fore.RED
    GREEN = Fore.GREEN
    YELLOW = Fore.YELLOW
    WHITE = Fore.WHITE
    MAGENTA = Fore.MAGENTA
    LIGHT_ORANGE = "\033[38;5;214m"
    ORANGE = "\033[38;5;208m"
    RESET = Style.RESET_ALL

def runner(day, path):
    document = docx.Document(path)
    table = document.tables[1]
    In_day = []
    for i in range(3, 7):
        Subject = str(table.cell(i, day).text.strip())
        if Subject != "—":
            In_day.append(Subject)

        out = f"Thứ {day}: {" - ".join(In_day)}"

    if str(table.cell(9, day).text.strip()) != "—":
       subj = str(table.cell(9, day).text.strip())
       out = out.__add__(f" | {subj.replace("\n", " ")}")

    if str(table.cell(10, day).text.strip()) != "—" and "|" in out:
       subj = str(table.cell(10, day).text.strip())
       out = out.__add__(f" - {subj.replace("\n", " ")}")
    elif str(table.cell(10, day).text.strip()) != "—" and "|" not in out:
        subj = str(table.cell(10, day).text.strip())
        out = out.__add__(f" | {subj.replace("\n", " ")}")


    print(out)
    return out

def start(path):
    copy = ""
    for day in range(2, 8):
        copy += runner(day, path)
        if day != 7:
            copy += "\n\n"
    return copy
    print(f"{line.LIGHT_ORANGE}{line.BOLD}Đã sao chép TKB rút gọn vào bảng nhớ tạm!{line.RESET}\n")